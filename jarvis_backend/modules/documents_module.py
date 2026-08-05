import os
import re
import shutil
import logging
from pathlib import Path
from typing import Optional, List, Dict
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)

try:
    import docx
except Exception:
    docx = None

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None


MAX_READ_CHARS = 32000


@dataclass
class DocumentResult:
    success: bool
    content: Optional[str] = None
    error: Optional[str] = None
    truncated: bool = False
    filename: Optional[str] = None


class DocumentsModule:
    ALLOWED_EXTENSIONS = {".txt", ".md", ".csv", ".docx", ".pdf"}
    ALLOWED_WRITE_EXTENSIONS = {".txt", ".md", ".docx"}

    def __init__(self, root: Optional[str] = None):
        raw_root = root or os.getenv("DOCUMENTS_ROOT", os.path.join(os.path.expanduser("~"), "Documents", "jarvis"))
        self.root = os.path.realpath(os.path.abspath(os.path.expanduser(raw_root)))
        os.makedirs(self.root, exist_ok=True)

    def _resolve(self, filename: str) -> str:
        path = os.path.realpath(os.path.join(self.root, filename))
        if not path.startswith(self.root + os.sep) and path != self.root:
            raise ValueError("Path escapes the documents root")
        return path

    def _token_budget_truncate(self, text: str) -> tuple[str, bool]:
        if len(text) <= MAX_READ_CHARS:
            return text, False
        return text[:MAX_READ_CHARS], True

    def _extract_text(self, path: str) -> str:
        ext = os.path.splitext(path)[1].lower()
        if ext == ".txt" or ext == ".md" or ext == ".csv":
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        if ext == ".docx":
            if docx is None:
                raise RuntimeError("python-docx is not installed")
            doc = docx.Document(path)
            return "\n".join(p.text for p in doc.paragraphs if p.text)
        if ext == ".pdf":
            if PdfReader is None:
                raise RuntimeError("pypdf is not installed")
            reader = PdfReader(path)
            parts = []
            for page in reader.pages:
                try:
                    parts.append(page.extract_text() or "")
                except Exception:
                    pass
            return "\n".join(parts)
        raise ValueError(f"Unsupported file type: {ext}")

    def list_documents(self, query: Optional[str] = None) -> List[str]:
        try:
            entries = []
            for name in os.listdir(self.root):
                if query and query.lower() not in name.lower():
                    continue
                entries.append(name)
            return sorted(entries)
        except Exception as e:
            logger.error(f"Failed to list documents: {e}")
            return []

    def read_document(self, filename: str) -> DocumentResult:
        path = self._resolve(filename)
        if not os.path.isfile(path):
            return DocumentResult(success=False, error="File not found", filename=filename)
        ext = os.path.splitext(path)[1].lower()
        if ext not in self.ALLOWED_EXTENSIONS:
            return DocumentResult(success=False, error=f"Unsupported file type: {ext}", filename=filename)
        try:
            text = self._extract_text(path)
            text, truncated = self._token_budget_truncate(text)
            note = " (truncated to ~8000 tokens)" if truncated else ""
            return DocumentResult(success=True, content=text + note, truncated=truncated, filename=filename)
        except Exception as e:
            logger.error(f"Failed to read document {filename}: {e}")
            return DocumentResult(success=False, error=str(e), filename=filename)

    def create_document(self, filename: str, content: str = "", overwrite: bool = False) -> DocumentResult:
        path = self._resolve(filename)
        ext = os.path.splitext(path)[1].lower()
        if ext not in self.ALLOWED_WRITE_EXTENSIONS:
            return DocumentResult(success=False, error=f"Unsupported file type for creation: {ext}")
        if os.path.exists(path) and not overwrite:
            return DocumentResult(success=False, error=f"File already exists: {filename}. Set overwrite to true to replace it.")
        try:
            os.makedirs(os.path.dirname(path), exist_ok=True)
            if ext == ".docx":
                if docx is None:
                    return DocumentResult(success=False, error="python-docx is not installed")
                doc = docx.Document()
                doc.add_paragraph(content)
                doc.save(path)
            else:
                with open(path, "w", encoding="utf-8") as f:
                    f.write(content)
            return DocumentResult(success=True, filename=filename)
        except Exception as e:
            logger.error(f"Failed to create document {filename}: {e}")
            return DocumentResult(success=False, error=str(e))

    def open_document(self, filename: str) -> DocumentResult:
        path = self._resolve(filename)
        if not os.path.isfile(path):
            return DocumentResult(success=False, error="File not found", filename=filename)
        try:
            if shutil.which("xdg-open"):
                import subprocess
                subprocess.Popen(["xdg-open", path])
            else:
                import subprocess
                subprocess.Popen(["open", path])
            return DocumentResult(success=True, filename=filename)
        except Exception as e:
            logger.error(f"Failed to open document {filename}: {e}")
            return DocumentResult(success=False, error=str(e), filename=filename)

    def append_document(self, filename: str, content: str) -> DocumentResult:
        path = self._resolve(filename)
        if not os.path.isfile(path):
            return DocumentResult(success=False, error="File not found", filename=filename)
        ext = os.path.splitext(path)[1].lower()
        if ext not in self.ALLOWED_WRITE_EXTENSIONS:
            return DocumentResult(success=False, error=f"Unsupported file type for editing: {ext}")
        try:
            if ext == ".docx":
                if docx is None:
                    return DocumentResult(success=False, error="python-docx is not installed")
                doc = docx.Document(path)
                doc.add_paragraph(content)
                doc.save(path)
            else:
                with open(path, "a", encoding="utf-8") as f:
                    f.write(content)
            return DocumentResult(success=True, filename=filename)
        except Exception as e:
            logger.error(f"Failed to append to document {filename}: {e}")
            return DocumentResult(success=False, error=str(e), filename=filename)

    def replace_document(self, filename: str, content: str) -> DocumentResult:
        path = self._resolve(filename)
        if not os.path.isfile(path):
            return DocumentResult(success=False, error="File not found", filename=filename)
        ext = os.path.splitext(path)[1].lower()
        if ext not in self.ALLOWED_WRITE_EXTENSIONS:
            return DocumentResult(success=False, error=f"Unsupported file type for editing: {ext}")
        try:
            if ext == ".docx":
                if docx is None:
                    return DocumentResult(success=False, error="python-docx is not installed")
                doc = docx.Document()
                doc.add_paragraph(content)
                doc.save(path)
            else:
                with open(path, "w", encoding="utf-8") as f:
                    f.write(content)
            return DocumentResult(success=True, filename=filename)
        except Exception as e:
            logger.error(f"Failed to replace document {filename}: {e}")
            return DocumentResult(success=False, error=str(e), filename=filename)

    def delete_document(self, filename: str) -> DocumentResult:
        path = self._resolve(filename)
        if not os.path.isfile(path):
            return DocumentResult(success=False, error="File not found", filename=filename)
        try:
            os.remove(path)
            return DocumentResult(success=True, filename=filename)
        except Exception as e:
            logger.error(f"Failed to delete document {filename}: {e}")
            return DocumentResult(success=False, error=str(e), filename=filename)
